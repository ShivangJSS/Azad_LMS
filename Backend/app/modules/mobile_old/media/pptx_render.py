"""
Renders PowerPoint slides to PNG so the app can show them inline.

There is no LibreOffice on this machine to convert to PDF, and Flutter cannot
render pptx. These decks are mostly positioned images, so each slide is
recomposed here from its shapes: pictures pasted at their recorded position
and size, then text drawn with its own size, colour and weight.

Rendered slides are cached on disk and reused, so a deck is only composed
once.
"""

import hashlib
import io
from pathlib import Path
from typing import Optional

CACHE_DIR = Path("uploads/_pptx_cache").resolve()

# Long edge of a rendered slide, in pixels.
TARGET_LONG_EDGE = 1600

# Windows ships these; the first one found is used for text.
FONT_CANDIDATES = [
    r"C:\Windows\Fonts\segoeui.ttf",
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\calibri.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]

BOLD_CANDIDATES = [
    r"C:\Windows\Fonts\segoeuib.ttf",
    r"C:\Windows\Fonts\arialbd.ttf",
    r"C:\Windows\Fonts\calibrib.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]

_PICTURE = 13


def _font_path(bold: bool) -> Optional[str]:
    for candidate in (BOLD_CANDIDATES if bold else FONT_CANDIDATES):
        if Path(candidate).exists():
            return candidate

    return None


def _load_font(size: int, bold: bool):
    from PIL import ImageFont

    path = _font_path(bold)

    if path is None:
        return ImageFont.load_default()

    try:
        return ImageFont.truetype(path, max(8, size))
    except OSError:
        return ImageFont.load_default()


def cache_key(source: Path) -> str:
    stat = source.stat()

    raw = f"{source}|{stat.st_size}|{int(stat.st_mtime)}"

    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def slide_dir(source: Path) -> Path:
    return CACHE_DIR / cache_key(source)


def render_deck(source: Path) -> int:
    """
    Render every slide once and return how many there are.
    """

    from pptx import Presentation

    target = slide_dir(source)

    done = target / "done"

    if done.exists():
        return int(done.read_text().strip() or 0)

    target.mkdir(parents=True, exist_ok=True)

    presentation = Presentation(str(source))

    width = presentation.slide_width or 1
    height = presentation.slide_height or 1

    scale = TARGET_LONG_EDGE / max(width, height)

    size = (max(1, int(width * scale)), max(1, int(height * scale)))

    slides = list(presentation.slides)

    for index, slide in enumerate(slides):
        image = _render_slide(slide, size, scale)
        image.save(target / f"{index + 1}.png", optimize=True)

    done.write_text(str(len(slides)))

    return len(slides)


def _render_slide(slide, size, scale):
    from PIL import Image, ImageDraw

    canvas = Image.new("RGB", size, "white")

    # Pictures first so text is never hidden behind them.
    for shape in slide.shapes:
        _paste_picture(canvas, shape, scale)

    draw = ImageDraw.Draw(canvas)

    for shape in slide.shapes:
        _draw_text(draw, shape, scale, size)

    return canvas


def _paste_picture(canvas, shape, scale) -> None:
    from PIL import Image

    if getattr(shape, "shape_type", None) != _PICTURE:
        return

    try:
        blob = shape.image.blob
    except Exception:
        return

    try:
        picture = Image.open(io.BytesIO(blob)).convert("RGBA")
    except Exception:
        return

    width = max(1, int((shape.width or 0) * scale))
    height = max(1, int((shape.height or 0) * scale))

    try:
        picture = picture.resize((width, height), Image.LANCZOS)
    except Exception:
        return

    canvas.paste(
        picture,
        (int((shape.left or 0) * scale), int((shape.top or 0) * scale)),
        picture,
    )


def _draw_text(draw, shape, scale, size) -> None:
    if not getattr(shape, "has_text_frame", False):
        return

    frame = shape.text_frame

    if not frame.text.strip():
        return

    left = int((shape.left or 0) * scale)
    top = int((shape.top or 0) * scale)

    box_width = max(40, int((shape.width or 0) * scale))

    y = top

    for paragraph in frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip()

        if not text:
            y += 8
            continue

        run = paragraph.runs[0]

        # Points to pixels: EMU scale is per-inch, and a point is 1/72 inch.
        points = run.font.size.pt if run.font.size else 18
        pixels = int(points * scale * 914400 / 72)

        font = _load_font(pixels, bool(run.font.bold))

        colour = _run_colour(run)

        for line in _wrap(draw, text, font, box_width):
            draw.text((left, y), line, font=font, fill=colour)
            y += int(pixels * 1.25)

        if y > size[1]:
            return


def _run_colour(run):
    try:
        rgb = run.font.color.rgb

        if rgb is not None:
            return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass

    return (33, 33, 33)


def _wrap(draw, text, font, max_width) -> list[str]:
    words = text.split()

    if not words:
        return []

    lines = []
    current = words[0]

    for word in words[1:]:
        trial = f"{current} {word}"

        if draw.textlength(trial, font=font) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word

    lines.append(current)

    return lines
